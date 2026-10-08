import os
import re
import json
from typing import Dict, Any

import streamlit as st
from pypdf import PdfReader
from docx import Document
from google import genai
from google.genai import types

st.set_page_config(page_title="ATS Resume Analyzer", page_icon="📄", layout="wide")


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(uploaded_file)
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")

    return re.sub(r"\n{3,}", "\n\n", text).strip()


def heuristic_score(text: str) -> int:
    lower = text.lower()
    score = 0

    checks = [
        ("email", r"[\w.+-]+@[\w-]+\.[\w.-]+", 10),
        ("phone", r"(?:\+?\d[\d\s().-]{7,}\d)", 8),
        ("experience", r"experience|work history|employment", 12),
        ("education", r"education|academic|degree|bachelor|master", 12),
        ("skills", r"skills|technical skills|competencies", 12),
        ("projects", r"projects|portfolio", 8),
        ("achievements", r"achievement|awards|certifications|certification", 8),
        (
            "action verbs",
            r"\b(led|built|developed|created|managed|improved|designed|implemented|optimized|analyzed)\b",
            10,
        ),
    ]

    for _, pattern, points in checks:
        if re.search(pattern, lower):
            score += points

    if 350 <= len(text) <= 7000:
        score += 10

    return min(score, 80)


def get_api_key() -> str | None:
    # Streamlit Cloud secrets are preferred; environment variables also work locally.
    try:
        secret_key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        secret_key = None

    return secret_key or os.getenv("GEMINI_API_KEY")


def analyze_with_gemini(text: str) -> Dict[str, Any]:
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY is not configured.")

    client = genai.Client(api_key=api_key)

    prompt = f"""
You are an ATS resume reviewer. Analyze the resume text below.

Return ONLY valid JSON using this schema:
{{
  "ats_score": integer 0-100,
  "summary": string,
  "strengths": [string],
  "improvements": [string],
  "missing_keywords": [string],
  "formatting_issues": [string],
  "action_items": [string]
}}

Be practical and honest. Do not invent experience. Judge ATS readability,
section structure, keyword coverage, measurable achievements, clarity, and
formatting based only on the supplied text.

RESUME:
{text[:30000]}
"""

    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )

    result = json.loads(response.text.strip())
    result["ats_score"] = max(
        0, min(100, int(result.get("ats_score", heuristic_score(text))))
    )
    return result


def display_results(result: Dict[str, Any]):
    score = result.get("ats_score", 0)

    st.subheader("ATS Score")
    st.progress(score / 100)
    st.metric("Score", f"{score}/100")
    st.write(result.get("summary", ""))

    col1, col2 = st.columns(2)

    with col1:
        st.markdown("### Strengths")
        for item in result.get("strengths", []):
            st.success(item)

        st.markdown("### Missing Keywords")
        for item in result.get("missing_keywords", []):
            st.warning(item)

    with col2:
        st.markdown("### Improvements")
        for item in result.get("improvements", []):
            st.info(item)

        st.markdown("### Formatting Issues")
        for item in result.get("formatting_issues", []):
            st.warning(item)

    st.markdown("### Action Plan")
    for i, item in enumerate(result.get("action_items", []), 1):
        st.write(f"**{i}.** {item}")


st.title("📄 ATS Resume Analyzer")
st.caption(
    "Upload a resume to get an ATS-focused score and practical improvement suggestions."
)

with st.sidebar:
    st.header("Settings")
    st.write("Supported formats: PDF, DOCX, TXT")
    st.info("Set GEMINI_API_KEY in Streamlit Secrets before deploying.")

uploaded = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt"],
)

if uploaded:
    try:
        resume_text = extract_text(uploaded)

        if not resume_text:
            st.error(
                "No readable text was found in this file. "
                "A scanned PDF may need OCR."
            )
        else:
            with st.expander("Preview extracted text"):
                st.text(resume_text[:12000])

            if st.button(
                "Analyze Resume",
                type="primary",
                use_container_width=True,
            ):
                with st.spinner("Analyzing your resume..."):
                    try:
                        result = analyze_with_gemini(resume_text)
                    except Exception as exc:
                        st.warning(
                            "Gemini analysis could not be completed, so a basic "
                            "local ATS check is shown instead."
                        )
                        result = {
                            "ats_score": heuristic_score(resume_text),
                            "summary": (
                                f"Gemini was unavailable: {exc}. "
                                "This fallback checks common ATS-friendly "
                                "resume elements."
                            ),
                            "strengths": [
                                "Resume text was successfully extracted."
                            ],
                            "improvements": [
                                "Configure GEMINI_API_KEY to get "
                                "AI-powered recommendations."
                            ],
                            "missing_keywords": [],
                            "formatting_issues": [],
                            "action_items": [
                                "Add measurable achievements and "
                                "role-specific keywords.",
                                "Use clear standard section headings.",
                            ],
                        }

                display_results(result)

    except Exception as exc:
        st.error(f"Could not read this file: {exc}")
else:
    st.info("Upload a resume to begin.")
